"""Read-only validation of the report directory and its referenced assets."""

from html.parser import HTMLParser
import json
from pathlib import Path
import re
from urllib.parse import unquote, urlsplit
import zipfile


class _ReportLinks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.ids, self.links, self.text = [], [], []
        self.hidden_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden_depth += 1
        attrs = dict(attrs)
        if attrs.get("id"):
            self.ids.append(attrs["id"])
        self.links.extend(attrs[key] for key in ("src", "href", "poster") if attrs.get(key))

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden_depth = max(0, self.hidden_depth - 1)

    def handle_data(self, data):
        if not self.hidden_depth:
            self.text.append(data)


def verify_report(output_dir):
    from yuntu_report import CURRENT_REPORT_SECTIONS

    output = Path(output_dir).resolve()
    errors, warnings, sizes = [], [], {}
    required = ("final-report.html", "final-report.md", "final-report-data.json", "presentation-report.json")
    for name in required:
        path = output / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"missing or empty report file: {name}")
    data = {}
    for name in required[2:]:
        try:
            value = json.loads((output / name).read_text(encoding="utf-8"))
            if not isinstance(value, dict):
                raise ValueError("expected JSON object")
            if name == "final-report-data.json":
                data = value
        except (OSError, ValueError) as exc:
            errors.append(f"invalid {name}: {exc}")
    parser = _ReportLinks()
    try:
        parser.feed((output / "final-report.html").read_text(encoding="utf-8"))
        for _, section_id, _ in CURRENT_REPORT_SECTIONS:
            if parser.ids.count(section_id) != 1:
                errors.append(f"missing or duplicate chapter anchor: {section_id}")
        # JavaScript legitimately contains undefined and braces; inspect visible text only.
        if re.search(r"\{\{[^{}]+\}\}|\b(?:TODO|undefined|NaN)\b", " ".join(parser.text)):
            errors.append("visible report text contains an unresolved placeholder")
    except (OSError, ValueError) as exc:
        errors.append(f"cannot read HTML: {exc}")
    local_refs = set()
    for reference in parser.links:
        parsed = urlsplit(reference)
        if parsed.scheme or parsed.netloc:
            continue
        if not parsed.path:
            if parsed.fragment and unquote(parsed.fragment) not in parser.ids:
                errors.append(f"broken navigation anchor: {reference}")
            continue
        local_refs.add(reference)
        resolved = (output / unquote(parsed.path)).resolve()
        if not resolved.is_relative_to(output) or not resolved.is_file():
            errors.append(f"missing or outside report local reference: {reference}")
    frames = data.get("highlightFrames", {})
    workbook = output / "highlight-frames.xlsx"
    if frames or workbook.exists():
        loaded = None
        try:
            with zipfile.ZipFile(workbook) as archive:
                if "xl/workbook.xml" not in archive.namelist() or archive.testzip() is not None:
                    raise ValueError("invalid workbook structure or CRC")
            from openpyxl import load_workbook
            loaded = load_workbook(workbook)
            for name in ("高光帧资产", "标准分类", "本批次标准"):
                if name not in loaded.sheetnames:
                    raise ValueError(f"missing workbook sheet: {name}")
            expected_images = sum(1 for frame in frames.values() if frame.get("available", True))
            if len(loaded["高光帧资产"]._images) != expected_images:
                raise ValueError(f"embedded asset images differ from successful frames: expected {expected_images}")
        except (OSError, ValueError, KeyError, TypeError, ImportError, zipfile.BadZipFile) as exc:
            errors.append(f"invalid highlight-frames.xlsx: {exc}")
        finally:
            if loaded is not None:
                loaded.close()
    actual = {p.relative_to(output).as_posix(): p for p in output.rglob("*") if p.is_file()}
    for name, path in actual.items():
        if path.is_symlink() or not path.resolve().is_relative_to(output):
            errors.append(f"report asset must be a local regular file: {name}")
        sizes[name] = path.stat().st_size
    report_bytes = sum(sizes.values())
    if report_bytes > 200 * 1024 * 1024:
        warnings.append("报告超过 200 MiB：交付 HTML 报告入口，并保留完整报告目录及相对路径；图片、视频和 Excel 依赖文件不能遗漏。")
    return {
        "valid": not errors, "errors": errors, "warnings": warnings,
        "outputDir": str(output), "reportPath": str(output / "final-report.html"),
        "localReferenceCount": len(local_refs), "reportFileCount": len(actual),
        "reportBytes": report_bytes,
        "largestFiles": [{"path": name, "bytes": size} for name, size in sorted(sizes.items(), key=lambda x: -x[1])[:8]],
    }
