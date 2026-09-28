"""Build report-aligned rows for the Chapter 04 highlight asset workbook."""

from collections import OrderedDict
from pathlib import Path
import sys
import hashlib
import posixpath
import tempfile
import zipfile
from xml.etree import ElementTree

from yuntu_top_materials import classify_brand


HIGHLIGHT_CATEGORY_LABELS = {
    "portrait": "人像画面",
    "product": "产品画面",
    "other": "其他画面",
}
OTHER_BRAND = "其他品牌"

IMAGE_MAX_WIDTH_PX = 140
IMAGE_MAX_HEIGHT_PX = 112
POINTS_PER_PIXEL = 0.75
GALLERY_COLUMNS = 4


def match_brands_from_title(title, brands):
    """Return every requested brand occurring in a material title."""
    result = classify_brand(title, brands=brands)
    values = result.get("matchedBrands") if isinstance(result, dict) else None
    if isinstance(values, list):
        return values
    value = result.get("value") if isinstance(result, dict) else None
    return [value] if isinstance(value, str) and value not in {"未知品牌", "歧义品牌"} else []


def match_brand_from_title(title, brands):
    """Backward-compatible single-label projection of title matches."""
    matches = match_brands_from_title(title, brands)
    return matches[0] if matches else OTHER_BRAND


def _chapter04_presentation(presentation):
    if not isinstance(presentation, dict):
        return {}
    chapters = presentation.get("chapters")
    if isinstance(chapters, dict) and isinstance(chapters.get("chapter04"), dict):
        return chapters["chapter04"]
    return presentation


def _mode_from_all_brand_scope(chapter, mode_id):
    brands = chapter.get("brands", []) if isinstance(chapter, dict) else []
    candidates = []
    for brand in brands:
        if not isinstance(brand, dict):
            continue
        modes = brand.get("classificationModes")
        if not isinstance(modes, list):
            continue
        mode = next((item for item in modes if isinstance(item, dict) and item.get("id") == mode_id), None)
        if mode is None:
            continue
        candidates.append((brand, mode))
    for brand, mode in candidates:
        if brand.get("id") == "all-brands" or brand.get("brand") == "全品牌":
            return mode
    return candidates[0][1] if candidates else None


def _category_map(categories, *, standard=False):
    mapping = {}
    order = []
    for category in categories if isinstance(categories, list) else []:
        if not isinstance(category, dict):
            continue
        category_id = str(category.get("id") or "")
        label = str(category.get("label") or "").strip()
        if standard:
            label = HIGHLIGHT_CATEGORY_LABELS.get(category_id, label or category_id)
        if not label:
            continue
        if label not in order:
            order.append(label)
        for frame in category.get("frames", []):
            if isinstance(frame, dict) and frame.get("frameRef"):
                mapping[str(frame["frameRef"])] = label
    return mapping, order


def _classification_maps(presentation, available_refs):
    chapter = _chapter04_presentation(presentation)
    standard_mode = _mode_from_all_brand_scope(chapter, "standard")
    batch_mode = _mode_from_all_brand_scope(chapter, "batch")
    standard_categories = standard_mode.get("categories") if isinstance(standard_mode, dict) else None
    if not isinstance(standard_categories, list):
        standard_categories = chapter.get("categories", []) if isinstance(chapter, dict) else []
    standard_map, standard_order = _category_map(standard_categories, standard=True)

    batch_reason = "未提供本批次标准分类"
    batch_map = {}
    batch_order = []
    mode_available = isinstance(batch_mode, dict) and batch_mode.get("available") is True
    if mode_available:
        batch_map, batch_order = _category_map(batch_mode.get("categories"), standard=False)
        if not batch_map:
            batch_reason = str(
                batch_mode.get("reason")
                or batch_mode.get("missingReason")
                or batch_mode.get("missing_reason")
                or "本批次标准分类没有有效分组"
            )
        elif not set(available_refs).issubset(batch_map):
            batch_map = {}
            batch_order = []
            batch_reason = "本批次标准分类未覆盖全部可用高光帧"
        else:
            batch_reason = ""
    elif isinstance(batch_mode, dict):
        batch_reason = str(
            batch_mode.get("reason")
            or batch_mode.get("missingReason")
            or batch_mode.get("missing_reason")
            or "本批次标准分类不可用"
        )
    return standard_map, standard_order, batch_map, batch_order, batch_reason


def build_highlight_export_rows(fact_pack, presentation=None):
    """Return one Chinese-keyed row per available frame, matching report cards."""
    facts = fact_pack.get("materialFacts", []) if isinstance(fact_pack, dict) else []
    brands = (fact_pack.get("filters", {}) or {}).get("brands", []) if isinstance(fact_pack, dict) else []
    available_refs = {
        str(frame.get("frame_ref"))
        for fact in facts
        if isinstance(fact, dict)
        for frame in fact.get("highlightFrames", [])
        if isinstance(frame, dict) and frame.get("available") is True and frame.get("frame_ref")
    }
    standard_map, standard_order, batch_map, batch_order, batch_reason = _classification_maps(
        presentation, available_refs
    )
    rows = []
    for fact in facts:
        if not isinstance(fact, dict):
            continue
        material_id = str(fact.get("material_id") or "")
        title = str(fact.get("title") or material_id)
        matches = match_brands_from_title(title, brands)
        brand = "、".join(matches) if matches else OTHER_BRAND
        for frame in fact.get("highlightFrames", []):
            if not isinstance(frame, dict) or frame.get("available") is not True:
                continue
            frame_ref = str(frame.get("frame_ref") or "")
            frame_type = frame.get("frame_type") or ("cover" if frame.get("is_cover") else "click_peak")
            peak_rank = frame.get("peak_rank")
            frame_label = "封面画面" if frame_type == "cover" else f"点击 Top{peak_rank} 波峰画面"
            rows.append(OrderedDict([
                ("品牌", brand),
                ("视觉类型", standard_map.get(frame_ref, "其他画面")),
                ("本批次标准分类", batch_map.get(frame_ref, "本批次标准不可用")),
                ("素材排名", fact.get("rank")),
                ("素材 ID", material_id),
                ("素材标题", title),
                ("画面类型", frame_label),
                ("秒数", frame.get("second")),
                ("点击峰值排名", peak_rank if peak_rank is not None else "-"),
                ("点击指数", frame.get("click_value") if frame.get("click_value") is not None else "-"),
                ("frame_ref", frame_ref),
                ("本地图片路径", frame.get("relative_path") or frame.get("image_path") or ""),
                ("__standard_category_order", standard_order.index(standard_map.get(frame_ref, "其他画面")) if standard_map.get(frame_ref, "其他画面") in standard_order else len(standard_order)),
                ("__batch_category_order", batch_order.index(batch_map[frame_ref]) if frame_ref in batch_map else len(batch_order)),
                ("__batch_classification_available", bool(batch_map)),
                ("__batch_missing_reason", batch_reason),
            ]))
    return sorted(
        rows,
        key=lambda row: (
            row.get("品牌", ""),
            row.get("__standard_category_order", 0),
            row.get("素材排名") if isinstance(row.get("素材排名"), (int, float)) else float("inf"),
            row.get("秒数") if isinstance(row.get("秒数"), (int, float)) else float("inf"),
            row.get("点击峰值排名") if isinstance(row.get("点击峰值排名"), (int, float)) else float("inf"),
        )
    )


def _resolve_frame_image_path(frame, asset_root=None):
    relative = str(frame.get("relative_path") or frame.get("image_path") or "").replace("\\", "/")
    if not relative:
        return ""
    path = Path(relative)
    candidates = [path] if path.is_absolute() else []
    if asset_root is not None and not path.is_absolute():
        root = Path(asset_root)
        candidates.append(root / relative)
        if relative.startswith("highlights/"):
            image_tail = Path(relative).relative_to("highlights")
            if root.name == "assets":
                candidates.append(root / "highlight-frames" / image_tail)
            candidates.extend((
                root / "assets" / "highlight-frames" / image_tail,
                root / "final-report" / "assets" / "highlight-frames" / image_tail,
            ))
        else:
            candidates.append(root / "final-report" / relative)
    return str(next((candidate.resolve() for candidate in candidates if candidate.is_file()), ""))


def export_highlight_workbook(fact_pack, presentation, output_path, asset_root=None):
    """Export report-aligned rows and a categorized image gallery to XLSX."""
    rows = build_highlight_export_rows(fact_pack, presentation)
    for row in rows:
        material_id = row.get("素材 ID")
        frame = next((frame for fact in fact_pack.get("materialFacts", []) if str(fact.get("material_id")) == str(material_id) for frame in fact.get("highlightFrames", []) if str(frame.get("frame_ref")) == str(row.get("frame_ref"))), {})
        row["__image_path"] = _resolve_frame_image_path(frame, asset_root)
    if not rows:
        return None
    output_path = Path(output_path).resolve()
    return _export_highlight_workbook_python(rows, output_path)


def _scaled_image_size(image_path, max_width=IMAGE_MAX_WIDTH_PX, max_height=IMAGE_MAX_HEIGHT_PX):
    from PIL import Image as PILImage

    with PILImage.open(image_path) as source:
        width, height = source.size
    scale = min(max_width / width, max_height / height, 1)
    return max(1, round(width * scale)), max(1, round(height * scale))


def _add_scaled_image(sheet, image_path, anchor, *, max_width=IMAGE_MAX_WIDTH_PX, max_height=IMAGE_MAX_HEIGHT_PX):
    from openpyxl.drawing.image import Image as XLImage

    width, height = _scaled_image_size(image_path, max_width, max_height)
    image = XLImage(image_path)
    image.width = width
    image.height = height
    sheet.add_image(image, anchor)
    return width, height


def _style_header(cells, *, fill_color="2C1810"):
    from openpyxl.styles import Alignment, Font, PatternFill

    for cell in cells:
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _build_asset_sheet(workbook, rows):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    headers = [
        "品牌", "视觉类型", "本批次标准分类", "素材排名", "素材 ID", "素材标题",
        "画面类型", "高光帧图片", "秒数", "点击峰值排名", "点击指数",
    ]
    sheet = workbook.active
    sheet.title = "高光帧资产"
    sheet.sheet_view.showGridLines = False
    sheet.append(headers)
    _style_header(sheet[1])
    sheet.row_dimensions[1].height = 30
    thin = Side(style="thin", color="DFCFBD")
    for row in rows:
        sheet.append([
            row.get("品牌"), row.get("视觉类型"), row.get("本批次标准分类"), row.get("素材排名"),
            row.get("素材 ID"), row.get("素材标题"), row.get("画面类型"), None,
            row.get("秒数"), row.get("点击峰值排名"), row.get("点击指数"),
        ])
        row_number = sheet.max_row
        for cell in sheet[row_number]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        sheet.cell(row_number, 5).number_format = "@"
        sheet.row_dimensions[row_number].height = IMAGE_MAX_HEIGHT_PX * POINTS_PER_PIXEL
        image_path = row.get("__image_path")
        if image_path and Path(image_path).is_file():
            try:
                _, image_height = _add_scaled_image(sheet, image_path, f"H{row_number}")
                sheet.row_dimensions[row_number].height = max(
                    sheet.row_dimensions[row_number].height or 0,
                    image_height * POINTS_PER_PIXEL,
                )
            except Exception as exc:
                print(f"[yuntu] 跳过无法读取的高光帧图片：{image_path}：{exc}", file=sys.stderr)
    widths = {
        "A": 18, "B": 14, "C": 20, "D": 10, "E": 22, "F": 42, "G": 22,
        "H": 20, "I": 12, "J": 12, "K": 12,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    end_row = max(2, len(rows) + 1)
    sheet.auto_filter.ref = f"A1:K{end_row}"
    table = Table(displayName="HighlightFramesTable", ref=f"A1:K{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    return sheet


def _build_gallery_sheet(workbook, title, rows, category_key, category_order, unavailable_reason=""):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.hyperlink import Hyperlink

    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    for column in range(1, GALLERY_COLUMNS + 1):
        sheet.column_dimensions[chr(64 + column)].width = 22
    sheet.freeze_panes = "A2"

    usable_rows = [row for row in rows if row.get("__image_path") and Path(row["__image_path"]).is_file()]
    brands = []
    for row in usable_rows:
        brand = row.get("品牌") or OTHER_BRAND
        if brand not in brands:
            brands.append(brand)
    category_order = list(category_order)
    for row in usable_rows:
        category = row.get(category_key)
        if category and category not in category_order:
            category_order.append(category)

    for column, brand in enumerate(brands, start=1):
        cell = sheet.cell(1, column, brand)
        cell.fill = PatternFill("solid", fgColor="2C1810")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28

    if not usable_rows:
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=GALLERY_COLUMNS)
        state = sheet.cell(3, 1, unavailable_reason or "暂无可用图片")
        state.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        state.font = Font(bold=True, color="2C1810")
        return sheet

    brand_rows = {}
    current_row = 3
    for brand in brands:
        brand_rows[brand] = current_row
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=GALLERY_COLUMNS)
        section = sheet.cell(current_row, 1, f"品牌：{brand}")
        section.fill = PatternFill("solid", fgColor="E9D8C4")
        section.font = Font(bold=True, color="2C1810")
        section.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1
        for category in category_order:
            scoped_rows = [row for row in usable_rows if row.get("品牌") == brand and row.get(category_key) == category]
            if not scoped_rows:
                continue
            sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=GALLERY_COLUMNS)
            category_cell = sheet.cell(current_row, 1, category)
            category_cell.fill = PatternFill("solid", fgColor="F4E8D8")
            category_cell.font = Font(bold=True, color="2C1810")
            category_cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1
            for offset in range(0, len(scoped_rows), GALLERY_COLUMNS):
                image_row = current_row
                id_row = current_row + 1
                sheet.row_dimensions[image_row].height = IMAGE_MAX_HEIGHT_PX * POINTS_PER_PIXEL
                sheet.row_dimensions[id_row].height = 24
                for column, row in enumerate(scoped_rows[offset:offset + GALLERY_COLUMNS], start=1):
                    image_path = row.get("__image_path")
                    try:
                        _, image_height = _add_scaled_image(
                            sheet, image_path, sheet.cell(image_row, column).coordinate
                        )
                        sheet.row_dimensions[image_row].height = max(
                            sheet.row_dimensions[image_row].height or 0,
                            image_height * POINTS_PER_PIXEL,
                        )
                    except Exception as exc:
                        print(f"[yuntu] 跳过无法读取的图片墙图片：{image_path}：{exc}", file=sys.stderr)
                    identifier = sheet.cell(id_row, column, f"素材 ID：{row.get('素材 ID') or '-'}")
                    identifier.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    identifier.font = Font(bold=True, color="2C1810")
                current_row += 3
    for column, brand in enumerate(brands, start=1):
        cell = sheet.cell(1, column)
        cell.hyperlink = Hyperlink(
            ref=cell.coordinate,
            location=f"'{title}'!A{brand_rows[brand]}",
        )
    return sheet


def _deduplicate_workbook_images(path):
    """Share byte-identical embedded media without resampling images or changing anchors."""
    with zipfile.ZipFile(path) as source:
        canonical, replacements = {}, {}
        for name in source.namelist():
            if not name.startswith("xl/media/") or name.endswith("/"):
                continue
            key = (Path(name).suffix.lower(), hashlib.sha256(source.read(name)).digest())
            if key in canonical:
                replacements[name] = canonical[key]
            else:
                canonical[key] = name
        if not replacements:
            return
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as stream:
            temporary = Path(stream.name)
        try:
            with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as target:
                for info in source.infolist():
                    if info.filename in replacements:
                        continue
                    data = source.read(info.filename)
                    if info.filename.endswith(".rels"):
                        document = ElementTree.fromstring(data)
                        changed = False
                        base = posixpath.dirname(posixpath.dirname(info.filename))
                        for relationship in document:
                            if relationship.get("TargetMode") == "External":
                                continue
                            value = relationship.get("Target", "")
                            resolved = (value.lstrip("/") if value.startswith("/")
                                        else posixpath.normpath(posixpath.join(base, value)))
                            if resolved in replacements:
                                relationship.set("Target", "/" + replacements[resolved])
                                changed = True
                        if changed:
                            data = ElementTree.tostring(document, encoding="utf-8", xml_declaration=True)
                    target.writestr(info, data)
            temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)


def _export_highlight_workbook_python(rows, output_path):
    """Generate the workbook with openpyxl and Pillow."""
    try:
        from openpyxl import Workbook
        from PIL import Image as PILImage  # noqa: F401 - validate Pillow before workbook authoring
    except Exception as exc:
        print(f"[yuntu] 高光帧资产导出不可用：Python XLSX 依赖加载失败：{exc}", file=sys.stderr)
        return None
    workbook = Workbook()
    _build_asset_sheet(workbook, rows)
    standard_order = ["人像画面", "产品画面", "其他画面"]
    batch_categories = {}
    batch_reason = ""
    for row in rows:
        label = row.get("本批次标准分类")
        if row.get("__batch_classification_available") and label and label != "本批次标准不可用":
            batch_categories.setdefault(label, row.get("__batch_category_order", len(batch_categories)))
        elif row.get("__batch_missing_reason") and not batch_reason:
            batch_reason = row["__batch_missing_reason"]
    batch_order = [label for label, _index in sorted(batch_categories.items(), key=lambda item: item[1])]
    standard_rows = [row for row in rows if row.get("__image_path")]
    batch_rows = [row for row in standard_rows if row.get("__batch_classification_available")]
    _build_gallery_sheet(workbook, "标准分类", standard_rows, "视觉类型", standard_order)
    _build_gallery_sheet(
        workbook,
        "本批次标准",
        batch_rows,
        "本批次标准分类",
        batch_order,
        unavailable_reason=f"本批次标准不可用：{batch_reason or '没有有效分类分组'}",
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix=".xlsx", delete=False) as stream:
        temporary = Path(stream.name)
    try:
        workbook.save(temporary)
        _deduplicate_workbook_images(temporary)
        temporary.replace(output_path)
    except Exception as exc:
        print(f"[yuntu] 高光帧资产导出不可用：Python XLSX 写入失败：{exc}", file=sys.stderr)
        return None
    finally:
        temporary.unlink(missing_ok=True)
    return output_path if output_path.is_file() and output_path.stat().st_size > 0 else None
