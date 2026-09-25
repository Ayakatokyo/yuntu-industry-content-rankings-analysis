import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorkbookImage

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import rpa_connector
from yuntu_report import _archive_report_directory


class DeliveryTests(unittest.TestCase):
    def fixture(self, root):
        output = root / "final-report"
        output.mkdir()
        (output / "assets").mkdir()
        Image.new("RGB", (8, 8), "red").save(output / "assets/frame.jpg")
        sections = ["top-material-combination", "audience-insights", "script-insights",
                    "highlight-insights", "material-deep-dive"]
        html = '<script>if (x !== undefined) {}</script>' + ''.join(
            f'<a href="#{key}">go</a><section id="{key}">正文</section>' for key in sections
        ) + '<img src="assets/frame.jpg"><a href="highlight-frames.xlsx">Excel</a>'
        (output / "final-report.html").write_text(html)
        (output / "final-report.md").write_text("# 完整报告")
        (output / "presentation-report.json").write_text('{"chapters": {}}')
        (output / "final-report-data.json").write_text(json.dumps({
            "status": "complete", "analysis": {}, "highlightFrames": {"f": {"available": True}}
        }))
        book = Workbook()
        book.active.title = "高光帧资产"
        book.active.add_image(WorkbookImage(output / "assets/frame.jpg"), "A1")
        book.create_sheet("标准分类")
        book.create_sheet("本批次标准")
        book.save(output / "highlight-frames.xlsx")
        _archive_report_directory(output)
        return output

    def verify(self, output):
        with contextlib.redirect_stdout(io.StringIO()) as stdout:
            code = rpa_connector.main(["verify-report", "--output-dir", str(output)])
        return code, json.loads(stdout.getvalue())

    def test_verifies_files_links_workbook_and_zip_without_unzip_or_jq(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = self.fixture(Path(tmp))
            code, result = self.verify(output)
            self.assertEqual(code, 0)
            self.assertTrue(result["valid"])
            self.assertEqual(result["errors"], [])
            self.assertEqual(result["localReferenceCount"], 2)
            self.assertEqual(result["archiveFileCount"], 6)

    def test_rejects_stale_zip_and_missing_asset_with_actionable_paths(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = self.fixture(Path(tmp))
            (output / "assets/frame.jpg").unlink()
            (output / "final-report.md").write_text("新报告")
            code, result = self.verify(output)
            self.assertEqual(code, 1)
            self.assertFalse(result["valid"])
            self.assertTrue(any("assets/frame.jpg" in message for message in result["errors"]))
            self.assertTrue(any("final-report.md" in message for message in result["errors"]))

    def test_zip_stores_compressed_media_but_deflates_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = self.fixture(Path(tmp))
            with zipfile.ZipFile(output.with_suffix(".zip")) as archive:
                self.assertEqual(archive.getinfo("final-report/assets/frame.jpg").compress_type, zipfile.ZIP_STORED)
                self.assertEqual(archive.getinfo("final-report/highlight-frames.xlsx").compress_type, zipfile.ZIP_STORED)
                self.assertEqual(archive.getinfo("final-report/final-report.html").compress_type, zipfile.ZIP_DEFLATED)

    def test_rejects_workbook_with_missing_embedded_frames_even_with_valid_crc(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = self.fixture(Path(tmp))
            book = Workbook()
            book.active.title = "高光帧资产"
            book.create_sheet("标准分类")
            book.create_sheet("本批次标准")
            book.save(output / "highlight-frames.xlsx")
            _archive_report_directory(output)
            code, result = self.verify(output)
            self.assertEqual(code, 1)
            self.assertTrue(any("embedded" in error for error in result["errors"]))


if __name__ == "__main__":
    unittest.main()
