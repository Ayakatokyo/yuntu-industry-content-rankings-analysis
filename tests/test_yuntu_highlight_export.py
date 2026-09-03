import sys
import unittest
from contextlib import redirect_stderr
from io import StringIO
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from yuntu_highlight_export import (  # noqa: E402
    build_highlight_export_rows,
    export_highlight_workbook,
)


class HighlightExportTests(unittest.TestCase):
    def _write_image(self, root, relative_path, size):
        path = root / relative_path
        path.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", size, color=(200, 120, 80)).save(path, format="JPEG")
        return path

    def _gallery_fixture(self, root):
        frames = []
        categories = []
        for index, (category, size) in enumerate(
            (("portrait", (100, 200)), ("product", (200, 100)), ("other", (120, 120))),
            start=1,
        ):
            ref = f"highlight:101:{category}:{index}"
            relative_path = f"highlights/101/{index:02d}.jpg"
            self._write_image(root, relative_path, size)
            frames.append({
                "frame_ref": ref,
                "second": index,
                "frame_type": "click_peak",
                "peak_rank": index,
                "click_value": 0.8,
                "available": True,
                "relative_path": relative_path,
            })
            categories.append({"id": category, "frames": [{"frameRef": ref}]})
        fact_pack = {
            "filters": {"brands": ["小仙炖"]},
            "materialFacts": [{
                "material_id": "101",
                "rank": 1,
                "title": "小仙炖测试素材",
                "highlightFrames": frames,
            }],
        }
        return fact_pack, {"categories": categories}

    def _dual_taxonomy_presentation(self, refs, *, batch_available=True):
        standard_categories = [
            {
                "id": category,
                "label": label,
                "frames": [{"frameRef": ref}],
            }
            for category, label, ref in (
                ("portrait", "人像画面", refs[0]),
                ("product", "产品画面", refs[1]),
                ("other", "其他画面", refs[2]),
            )
        ]
        batch_categories = [
            {
                "id": "人物与产品同框",
                "label": "人物与产品同框",
                "frames": [{"frameRef": refs[0]}],
            },
            {
                "id": "包装信息",
                "label": "包装信息",
                "frames": [{"frameRef": refs[1]}, {"frameRef": refs[2]}],
            },
        ]
        return {
            "categories": standard_categories,
            "brands": [{
                "brand": "全品牌",
                "classificationModes": [
                    {"id": "standard", "available": True, "categories": standard_categories},
                    {
                        "id": "batch",
                        "available": batch_available,
                        "reason": "整批画面理解证据不足" if not batch_available else "",
                        "categories": batch_categories if batch_available else [],
                    },
                ],
            }],
        }

    def _dual_gallery_fixture(self, root):
        fact_pack, _presentation = self._gallery_fixture(root)
        refs = [
            frame["frame_ref"]
            for frame in fact_pack["materialFacts"][0]["highlightFrames"]
        ]
        return fact_pack, self._dual_taxonomy_presentation(refs)

    def test_export_uses_openpyxl_and_pillow_by_default(self):
        fact_pack = {
            "filters": {"brands": ["小仙炖"]},
            "materialFacts": [{
                "material_id": "101",
                "title": "小仙炖测试素材",
                "highlightFrames": [{"frame_ref": "f1", "available": True}],
            }],
        }
        with tempfile.TemporaryDirectory() as tmp, redirect_stderr(StringIO()) as stderr:
            result = export_highlight_workbook(
                fact_pack, {}, Path(tmp) / "highlight-frames.xlsx"
            )
        self.assertIsNotNone(result)
        self.assertNotIn("高光帧资产导出不可用", stderr.getvalue())

    def test_export_writes_workbook_without_images(self):
        fact_pack = {
            "filters": {"brands": ["小仙炖"]},
            "materialFacts": [{
                "material_id": "101", "rank": 1, "title": "小仙炖测试素材",
                "highlightFrames": [{"frame_ref": "f1", "available": True, "second": 0}],
            }],
        }
        with tempfile.TemporaryDirectory() as tmp:
            output = export_highlight_workbook(fact_pack, {}, Path(tmp) / "highlight-frames.xlsx")
            self.assertTrue(output.is_file())
            self.assertGreater(output.stat().st_size, 0)

    def test_builds_report_aligned_rows_and_excludes_unavailable_frames(self):
        fact_pack = {
            "filters": {"brands": ["小仙炖"]},
            "materialFacts": [{
                "material_id": "101",
                "rank": 3,
                "title": "小仙炖测试素材",
                "highlightFrames": [
                    {"frame_ref": "highlight:101:cover:0", "second": 0, "frame_type": "cover", "available": True, "relative_path": "highlights/101/01-000s.jpg"},
                    {"frame_ref": "highlight:101:click_peak:2", "second": 2, "frame_type": "click_peak", "peak_rank": 1, "click_value": 0.8, "available": True, "relative_path": "highlights/101/02-002s.jpg"},
                    {"frame_ref": "highlight:101:click_peak:9", "second": 9, "frame_type": "click_peak", "available": False, "relative_path": ""},
                ],
            }],
        }
        rows = build_highlight_export_rows(fact_pack)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["品牌"], "小仙炖")
        self.assertEqual(rows[0]["画面类型"], "封面画面")
        self.assertEqual(rows[1]["画面类型"], "点击 Top1 波峰画面")
        self.assertEqual(rows[1]["秒数"], 2)
        self.assertEqual(rows[1]["点击指数"], 0.8)
        self.assertEqual(rows[1]["素材标题"], "小仙炖测试素材")

    def test_builds_rows_with_standard_and_batch_taxonomy_labels(self):
        fact_pack = {
            "filters": {"brands": ["小仙炖"]},
            "materialFacts": [{
                "material_id": "101",
                "rank": 1,
                "title": "小仙炖测试素材",
                "highlightFrames": [
                    {"frame_ref": "f1", "available": True, "second": 0},
                    {"frame_ref": "f2", "available": True, "second": 2},
                    {"frame_ref": "f3", "available": True, "second": 4},
                ],
            }],
        }
        presentation = self._dual_taxonomy_presentation(["f1", "f2", "f3"])

        rows = build_highlight_export_rows(fact_pack, presentation)

        self.assertEqual(
            [(row["视觉类型"], row["本批次标准分类"]) for row in rows],
            [("人像画面", "人物与产品同框"), ("产品画面", "包装信息"), ("其他画面", "包装信息")],
        )

    def test_marks_batch_taxonomy_unavailable_without_standard_fallback(self):
        fact_pack = {
            "filters": {"brands": ["小仙炖"]},
            "materialFacts": [{
                "material_id": "101",
                "rank": 1,
                "title": "小仙炖测试素材",
                "highlightFrames": [{"frame_ref": "f1", "available": True}],
            }],
        }
        presentation = self._dual_taxonomy_presentation(["f1", "f2", "f3"], batch_available=False)

        rows = build_highlight_export_rows(fact_pack, presentation)

        self.assertEqual(rows[0]["视觉类型"], "人像画面")
        self.assertEqual(rows[0]["本批次标准分类"], "本批次标准不可用")
        self.assertEqual(rows[0]["__batch_missing_reason"], "整批画面理解证据不足")

    def test_exports_highlight_images_with_source_ratio_and_covering_row_height(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fact_pack, presentation = self._gallery_fixture(root)
            output = export_highlight_workbook(
                fact_pack, presentation, root / "highlight-frames.xlsx", asset_root=root
            )
            workbook = load_workbook(output)
            sheet = workbook["高光帧资产"]
            self.assertEqual(len(sheet._images), 3)
            self.assertEqual(sheet.auto_filter.ref, "A1:K4")
            self.assertEqual(sheet.freeze_panes, "A2")
            expected_ratios = [100 / 200, 200 / 100, 1]
            for image, source_ratio in zip(sheet._images, expected_ratios):
                self.assertEqual(image.anchor._from.col, 7)
                row = image.anchor._from.row + 1
                width = image.anchor.ext.width / 9525
                height = image.anchor.ext.height / 9525
                self.assertGreaterEqual(sheet.row_dimensions[row].height, height * 0.75)
                self.assertAlmostEqual(width / height, source_ratio, places=2)
            self.assertGreaterEqual(sheet.column_dimensions["G"].width, 18)
            self.assertIsNotNone(sheet.tables["HighlightFramesTable"])

    def test_resolves_images_from_report_assets_root(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fact_pack, presentation = self._dual_gallery_fixture(root)
            asset_root = root / "assets"
            for frame in fact_pack["materialFacts"][0]["highlightFrames"]:
                source = root / frame["relative_path"]
                target = asset_root / "highlight-frames" / Path(frame["relative_path"]).relative_to("highlights")
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(source.read_bytes())
            output = export_highlight_workbook(
                fact_pack, presentation, root / "highlight-frames.xlsx", asset_root=asset_root
            )
            workbook = load_workbook(output)

        self.assertEqual(len(workbook["标准分类"]._images), 3)
        self.assertEqual(len(workbook["本批次标准"]._images), 3)
        workbook.close()

    def test_asset_sheet_contains_both_taxonomy_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fact_pack, presentation = self._dual_gallery_fixture(root)
            output = export_highlight_workbook(
                fact_pack, presentation, root / "highlight-frames.xlsx", asset_root=root
            )
            workbook = load_workbook(output)

        sheet = workbook["高光帧资产"]
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            ["品牌", "视觉类型", "本批次标准分类", "素材排名", "素材 ID", "素材标题", "画面类型", "高光帧图片", "秒数", "点击峰值排名", "点击指数"],
        )
        self.assertEqual(sheet.auto_filter.ref, "A1:K4")
        self.assertEqual(sheet._images[0].anchor._from.col, 7)
        workbook.close()

    def test_exports_highlight_gallery_with_all_categories_and_internal_links(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fact_pack, presentation = self._dual_gallery_fixture(root)
            output = export_highlight_workbook(
                fact_pack, presentation, root / "highlight-frames.xlsx", asset_root=root
            )
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, ["高光帧资产", "标准分类", "本批次标准"])
            for sheet_name, expected_categories in (
                ("标准分类", ("人像画面", "产品画面", "其他画面")),
                ("本批次标准", ("人物与产品同框", "包装信息")),
            ):
                sheet = workbook[sheet_name]
                text = "\n".join(
                    str(cell.value or "")
                    for row in sheet.iter_rows()
                    for cell in row
                )
                self.assertIn("小仙炖", text)
                for category in expected_categories:
                    self.assertIn(category, text)
                self.assertIn("素材 ID：101", text)
                self.assertNotIn("小仙炖测试素材", text)
                self.assertEqual(len(sheet._images), 3)
                self.assertGreaterEqual(sheet.column_dimensions[get_column_letter(1)].width, 18)
            self.assertTrue(all(cell.hyperlink.location for cell in workbook["标准分类"][1] if cell.value))
            self.assertTrue(all(cell.hyperlink.location for cell in workbook["本批次标准"][1] if cell.value))
            workbook.close()

    def test_unavailable_batch_gallery_does_not_fallback_to_standard_images(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fact_pack, _presentation = self._gallery_fixture(root)
            refs = [frame["frame_ref"] for frame in fact_pack["materialFacts"][0]["highlightFrames"]]
            presentation = self._dual_taxonomy_presentation(refs, batch_available=False)
            output = export_highlight_workbook(
                fact_pack, presentation, root / "highlight-frames.xlsx", asset_root=root
            )
            workbook = load_workbook(output)

        sheet = workbook["本批次标准"]
        text = "\n".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
        self.assertIn("本批次标准不可用", text)
        self.assertIn("整批画面理解证据不足", text)
        self.assertEqual(len(sheet._images), 0)
        self.assertNotIn("素材 ID：101", text)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
